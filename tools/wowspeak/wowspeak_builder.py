#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
WOW Speak Speaking-Club deck builder (v3).

Changes vs build_v2.py:
  * Wordwall is gone. The vocabulary game is built INSIDE the .pptx by
    wowspeak_games.py (cfg['games']).
  * The video link hangs on a button SHAPE (click_action.hyperlink.address),
    never on a text run - editors repaint linked runs blue+underlined and only
    the text stays clickable.
  * A local video file (cfg['video_file']) is embedded with add_movie.
  * No hard-coded theme emoji anywhere: title / farewell read cfg['emoji'].
  * save() writes exactly where you ask (prs.save(out)); no /mnt/... copying.

Public helpers (wowspeak_games.py builds on them):
    PASTELS, BLOB, DARK/HOTPINK/GOLD/WHITE/INK/RED/TEAL, TF/BF
    C, set_bg, blob, shadow, card, text, lnk, fit, logo, band, new_slide
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from PIL import Image

PASTELS={'pink':'FFE4F0','blue':'D4ECFF','yellow':'FFF7D4','lavender':'EDE4FA','mint':'D4F5E6','peach':'FFE6D0'}
DARK='2C2040'; HOTPINK='E91E8C'; GOLD='FFD700'; WHITE='FFFFFF'; INK='3A2E55'; RED='D32F2F'; TEAL='1E9E7A'
BLOB=['FFB6D9','A9D6FF','FFE89A','C9B3F0','9EE9C6','FFC79E']; TF='Trebuchet MS'; BF='Calibri'
SLIDE_W=13.3333; SLIDE_H=7.5

def C(h): return RGBColor.from_string(h)
def set_bg(s,h):
    f=s.background.fill; f.solid(); f.fore_color.rgb=C(h)
def _al(el,p): el.append(el.makeelement(qn('a:alpha'),{'val':str(int(p*1000))}))
def blob(s,cx,cy,sz,col,al=40):
    sh=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(cx-sz/2),Inches(cy-sz/2),Inches(sz),Inches(sz))
    sh.line.fill.background(); sh.fill.solid(); sh.fill.fore_color.rgb=C(col)
    g=sh.fill.fore_color._xFill.find(qn('a:srgbClr'))
    if g is not None:_al(g,al)
    sh.shadow.inherit=False
    return sh
def shadow(sh):
    sp=sh._element.spPr
    for e in sp.findall(qn('a:effectLst')): sp.remove(e)
    eff=sp.makeelement(qn('a:effectLst'),{}); o=eff.makeelement(qn('a:outerShdw'),{'blurRad':'90000','dist':'45000','dir':'5400000','rotWithShape':'0'})
    c=o.makeelement(qn('a:srgbClr'),{'val':'6A4E8C'}); c.append(c.makeelement(qn('a:alpha'),{'val':'30000'})); o.append(c); eff.append(o); sp.append(eff)
def card(s,x,y,w,h,fill=WHITE,rad=0.1,sh=True):
    r=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    try:r.adjustments[0]=rad
    except:pass
    r.fill.solid(); r.fill.fore_color.rgb=C(fill); r.line.fill.background(); r.shadow.inherit=False
    if sh: shadow(r)
    return r
def text(s,x,y,w,h,runs,align=PP_ALIGN.LEFT,anchor=MSO_ANCHOR.TOP,sa=6,ls=1.0):
    tb=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=tb.text_frame
    tf.word_wrap=True; tf.vertical_anchor=anchor
    for m in ('left','right','top','bottom'): setattr(tf,f'margin_{m}',Pt(2))
    first=True
    for para in runs:
        p=tf.paragraphs[0] if first else tf.add_paragraph(); first=False
        p.alignment=align; p.space_after=Pt(sa); p.space_before=Pt(0); p.line_spacing=ls
        for (t,sz,col,b,fn,*rest) in para:
            it=rest[0] if rest else False
            r=p.add_run(); r.text=t; r.font.size=Pt(sz); r.font.bold=b; r.font.italic=it; r.font.name=fn; r.font.color.rgb=C(col)
    return tb
def lnk(p,t,url,sz,col,fn=BF,b=False):
    r=p.add_run(); r.text=t; r.font.size=Pt(sz); r.font.bold=b; r.font.name=fn; r.font.color.rgb=C(col); r.hyperlink.address=url
def fit(s,path,x,y,w,h):
    """Place a picture inside the box (x,y,w,h) keeping its aspect ratio."""
    iw,ih=Image.open(path).size; ar=iw/ih; bar=w/h
    if ar>bar: dw=w; dh=w/ar
    else: dh=h; dw=h*ar
    return s.shapes.add_picture(path,Inches(x+(w-dw)/2),Inches(y+(h-dh)/2),Inches(dw),Inches(dh))
def logo(s,dark=False):
    col=GOLD if dark else HOTPINK; h='💛' if dark else '💜'
    tb=s.shapes.add_textbox(Inches(11.55),Inches(6.95),Inches(1.7),Inches(0.45)); tf=tb.text_frame; tf.word_wrap=False
    p=tf.paragraphs[0]; p.alignment=PP_ALIGN.RIGHT
    r=p.add_run(); r.text='WOW Speak '; r.font.size=Pt(14); r.font.bold=True; r.font.name=TF; r.font.color.rgb=C(col)
    r2=p.add_run(); r2.text=h; r2.font.size=Pt(14); r2.font.name=BF
def band(s,t,dark=False,y=0.45,size=32):
    col=GOLD if dark else HOTPINK
    if dark: text(s,0.7,y,11.9,1.0,[[(t,size,col,True,TF)]],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
    else:
        card(s,1.4,y,10.5,1.05,fill=WHITE,rad=0.5); text(s,1.5,y,10.3,1.05,[[(t,size,col,True,TF)]],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)

def new_presentation():
    prs=Presentation(); prs.slide_width=Emu(12192000); prs.slide_height=Emu(6858000); return prs
def new_slide(prs,bg):
    sl=prs.slides.add_slide(prs.slide_layouts[6]); set_bg(sl,bg); return sl

# Reusable Advanced discussion toolbox (same across all Advanced decks)
USEFUL=[('💬 Give your opinion','In my opinion… / I think… / If you ask me…'),
        ('👍 Agree','I agree because… / Good point! / Exactly!'),
        ('🤔 Disagree','I’m not sure… / I see your point, but… / On the other hand…'),
        ('💭 React & ask','Really? / That’s interesting! / Why do you think so?')]

GAME_NOTE=('Игры встроены в презентацию. Работают ТОЛЬКО в режиме показа (F5). '
           'Клик мимо кнопки НЕ листает слайд — дальше переходи стрелкой →. '
           'Подписи под картинками в играх убраны специально: иначе задание проверяет чтение, а не знание слова. '
           'Экраны «Correct / Try again» лежат в конце файла — листая урок стрелками, ты их не увидишь.')

def build(cfg,out):
    """Render one deck. cfg schema: see HANDOFF.md §6."""
    import wowspeak_games as G
    lvl=cfg['level']  # 'beg' or 'adv'
    prs=new_presentation()
    def ns(bg): return new_slide(prs,bg)
    T=cfg['theme']; EM=cfg['emoji']
    helpers=[]

    # S1 title
    s=ns(cfg['title_bg'])
    blob(s,0.4,0.6,3.2,BLOB[2]); blob(s,12.8,7.0,3.6,BLOB[1]); blob(s,12.6,0.3,2.4,BLOB[0],35)
    text(s,0.6,0.5,3.0,1.3,[[(EM[0],80,INK,False,BF)]],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
    tw_=min(10.4,3.2+len(T)*0.42); tx=(SLIDE_W-tw_)/2
    card(s,tx,2.2,tw_,1.9,fill=WHITE,rad=0.18)
    text(s,tx,2.3,tw_,1.5,[[(EM[1]+' ',34,HOTPINK,True,BF),(T,36,HOTPINK,True,TF),(' '+EM[2],34,HOTPINK,True,BF)]],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
    card(s,3.9,4.45,5.53,0.9,fill=WHITE,rad=0.5)
    text(s,4.0,4.45,5.33,0.9,[[('🗣️ ',20,INK,False,BF),(cfg['subtitle'],22,'7A5C9E',True,TF)]],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
    text(s,1.5,5.6,10.3,0.8,[[(' '.join(EM[3]),34,INK,False,BF)]],align=PP_ALIGN.CENTER)
    logo(s)

    # S2 teacher notes - the slide is reserved here, the text is written at the
    # very end, when the real slide numbers are known (rule 9).
    s_notes=ns(DARK)

    # S3 mood
    s=ns(PASTELS['blue'])
    blob(s,0.4,7.1,3.0,BLOB[4]); blob(s,13.0,0.4,2.8,BLOB[2])
    band(s,'😊 How are you today? 😊',size=34)
    text(s,1.0,1.65,11.3,0.5,[[(cfg['mood_q'],18,'7A5C9E',True,TF)]],align=PP_ALIGN.CENTER)
    moods=cfg['moods']
    n=5; gap=0.35; cw=(11.6-(n-1)*gap)/n; gx=(SLIDE_W-(n*cw+(n-1)*gap))/2; gy=2.5; ch=2.3
    for i,(e,l) in enumerate(moods):
        x=gx+i*(cw+gap); card(s,x,gy,cw,ch,fill=WHITE,rad=0.2)
        text(s,x,gy+0.2,cw,1.1,[[(e,46,INK,False,BF)]],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
        text(s,x,gy+1.4,cw,0.6,[[(l,16,HOTPINK,True,TF)]],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
    logo(s)

    # S4 LEAD-IN
    s=ns(PASTELS['yellow'])
    blob(s,0.3,0.4,2.8,BLOB[0]); blob(s,13.1,7.2,3.2,BLOB[3])
    band(s,'💬 Lead-in – let’s talk! '+EM[0],size=30)
    card(s,1.2,1.75,10.93,4.95,fill=WHITE,rad=0.06)
    runs=[[('🗨️ ',18,TEAL,True,BF),('Discuss with your teacher:',19,TEAL,True,TF)]]
    for q,e in cfg['lead_in']:
        runs.append([(e+'  ',15.5,INK,False,BF),(q,15.5,INK,False,BF)])
    text(s,1.6,2.05,10.1,4.4,runs,align=PP_ALIGN.LEFT,sa=8,ls=1.05)
    logo(s)

    # S5 NEW WORDS - the ONLY slide where the pictures carry captions
    s=ns(PASTELS['mint'])
    blob(s,0.4,0.5,3.0,BLOB[1]); blob(s,13.0,7.1,3.0,BLOB[0])
    band(s,cfg['vocab_title'],size=32)
    vocab=cfg['vocab']
    if lvl=='beg':
        cols=3; gap=0.35; cw=(11.4-(cols-1)*gap)/cols; gx=(SLIDE_W-(cols*cw+(cols-1)*gap))/2; gy=1.9; ch=2.3
        for i,(word,img) in enumerate(vocab):
            r=i//cols; c=i%cols; x=gx+c*(cw+gap); y=gy+r*(ch+0.2)
            card(s,x,y,cw,ch,fill=WHITE,rad=0.12)
            if img and os.path.exists(img): fit(s,img,x+0.15,y+0.12,cw-0.3,ch-0.75)
            text(s,x+0.05,y+ch-0.6,cw-0.1,0.5,[[(word,15,HOTPINK,True,TF)]],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
    else:
        cols=3; gap=0.35; cw=(11.7-(cols-1)*gap)/cols; gx=(SLIDE_W-(cols*cw+(cols-1)*gap))/2; gy=1.9; ch=2.3
        for i,(word,defn) in enumerate(vocab):
            r=i//cols; c=i%cols; x=gx+c*(cw+gap); y=gy+r*(ch+0.2)
            card(s,x,y,cw,ch,fill=WHITE,rad=0.1)
            text(s,x+0.1,y+0.2,cw-0.2,0.55,[[(word,18,HOTPINK,True,TF)]],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
            text(s,x+0.16,y+0.82,cw-0.32,ch-1.0,[[(defn,14,INK,False,BF)]],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.TOP,ls=1.02)
    logo(s)

    # S6 TALK WITH THE WORDS
    if lvl=='beg':
        s=ns(PASTELS['peach'])
        blob(s,0.4,0.5,3.0,BLOB[3]); blob(s,13.0,7.1,3.0,BLOB[5])
        mini=cfg['mini']
        band(s,mini['title'],size=30)
        card(s,1.4,1.8,10.5,4.9,fill=WHITE,rad=0.06)
        text(s,1.7,2.05,9.9,0.6,[[('🗣️ Say: ',17,TEAL,True,BF),(mini['frame'],17,TEAL,True,TF)]],align=PP_ALIGN.LEFT)
        items=mini['items']
        half=(len(items)+1)//2
        r1=[[(e+'  ',16,HOTPINK,True,BF),(t,16,INK,False,BF)] for t,e in items[:half]]
        r2=[[(e+'  ',16,HOTPINK,True,BF),(t,16,INK,False,BF)] for t,e in items[half:]]
        text(s,1.9,2.9,4.9,3.5,r1,align=PP_ALIGN.LEFT,sa=13,ls=1.05)
        text(s,7.0,2.9,4.7,3.5,r2,align=PP_ALIGN.LEFT,sa=13,ls=1.05)
        logo(s)
    else:
        s=ns(PASTELS['peach'])
        blob(s,0.4,0.5,3.0,BLOB[3]); blob(s,13.0,7.1,3.0,BLOB[5])
        band(s,'💬 Useful language for discussion 🗣️',size=28)
        gx=1.4; gy=1.95; cw=5.15; ch=2.15; gapx=0.55; gapy=0.35
        for i,(hd,ph) in enumerate(USEFUL):
            r=i//2; c=i%2; x=gx+c*(cw+gapx); y=gy+r*(ch+gapy)
            card(s,x,y,cw,ch,fill=WHITE,rad=0.1)
            text(s,x+0.25,y+0.2,cw-0.5,0.6,[[(hd,18,HOTPINK,True,TF)]],align=PP_ALIGN.LEFT,anchor=MSO_ANCHOR.MIDDLE)
            text(s,x+0.3,y+0.85,cw-0.6,ch-1.0,[[(ph,15,INK,False,BF)]],align=PP_ALIGN.LEFT,anchor=MSO_ANCHOR.TOP,ls=1.1)
        logo(s)

    # GAMES (built into the deck; no Wordwall)
    before={sl.slide_id for sl in prs.slides}
    for spec in cfg.get('games',[]):
        helpers+=G.play(prs,spec,cfg)
    hid={sl.slide_id for sl in helpers}
    game_slides=[sl for sl in prs.slides if sl.slide_id not in before and sl.slide_id not in hid]

    # video + discussion
    s=s_video=ns(DARK)
    blob(s,0.3,0.4,2.6,'5A4A7A',50); blob(s,13.0,7.2,3.0,'5A4A7A',50)
    band(s,'🎬 Watch the Video! 🍿',dark=True,size=34)
    vf=cfg.get('video_file')
    if vf and os.path.exists(vf):
        # local file embedded straight into the slide
        s.shapes.add_movie(vf,Inches(3.5),Inches(1.75),Inches(6.33),Inches(2.7),
                           poster_frame_image=cfg.get('video_poster'),mime_type='video/mp4')
    else:
        card(s,3.5,1.75,6.33,2.7,fill='000000',rad=0.05)
        text(s,3.6,1.75,6.13,2.7,[[('🎬',70,'6A5A8A',False,BF)]],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
    url=cfg.get('video_url')
    if url:
        # the link lives on the BUTTON SHAPE, not on a text run
        b=card(s,4.4,4.6,4.53,0.75,fill=GOLD,rad=0.5)
        text_in_shape(b,'▶️  Open the video',18,DARK,True,TF)
        b.click_action.hyperlink.address=url
    runs=[[('💬 After watching, discuss:',16,GOLD,True,TF)]]
    for q,e in cfg['video_qs']:
        runs.append([(e+'  ',14.5,'F2ECFF',False,BF),(q,14.5,'F2ECFF',False,BF)])
    text(s,2.3,5.5,8.73,1.5,runs,align=PP_ALIGN.LEFT,sa=6,ls=1.03)
    logo(s,dark=True)

    # MAIN SPEAKING
    if lvl=='beg':
        s=s_speak=ns(PASTELS['lavender'])
        blob(s,0.3,0.4,2.8,BLOB[4]); blob(s,13.1,7.2,3.0,BLOB[1])
        band(s,cfg['speak_title'],size=30)
        card(s,1.4,1.75,10.5,5.0,fill=WHITE,rad=0.06)
        text(s,1.8,2.0,9.8,0.5,[[('💡 ',15,'C2722A',True,BF),(cfg['frame_hint'],15,'C2722A',True,BF)]],align=PP_ALIGN.LEFT)
        runs=[[('• ',18,HOTPINK,True,BF),(fr,18,INK,False,BF)] for fr in cfg['frames']]
        text(s,2.0,2.7,9.4,3.9,runs,align=PP_ALIGN.LEFT,sa=14,ls=1.05)
        logo(s)
    else:
        s=s_speak=ns(PASTELS['lavender'])
        blob(s,0.3,0.4,2.8,BLOB[4]); blob(s,13.1,7.2,3.0,BLOB[1])
        d=cfg['disc']
        band(s,'🗣️ '+d['headline'],size=30)
        card(s,1.6,1.7,10.13,1.15,fill='FFF0F7',rad=0.18)
        text(s,1.8,1.7,9.73,1.15,[[(d['statement'],21,HOTPINK,True,TF)]],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
        yy=3.05
        if d.get('options'):
            opts=d['options']; n=len(opts); ow=min(3.2,(10.5-(n-1)*0.4)/n); tot=n*ow+(n-1)*0.4; ox=(SLIDE_W-tot)/2
            for i,op in enumerate(opts):
                x=ox+i*(ow+0.4); card(s,x,yy,ow,0.75,fill=WHITE,rad=0.4)
                text(s,x,yy,ow,0.75,[[(op,16,INK,True,BF)]],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
            yy+=1.0
        card(s,1.6,yy,10.13,6.55-yy,fill=WHITE,rad=0.06)
        runs=[[('📋 ',17,TEAL,True,BF),('Your task:',18,TEAL,True,TF)]]
        for i,tk in enumerate(d['tasks']):
            runs.append([(f'{i+1}. ',16,HOTPINK,True,BF),(tk,16,INK,False,BF)])
        text(s,1.95,yy+0.2,9.4,6.35-yy,runs,align=PP_ALIGN.LEFT,sa=10,ls=1.04)
        logo(s)

    # thank you
    s=ns(PASTELS['pink'])
    blob(s,0.4,0.5,3.6,BLOB[1]); blob(s,12.9,7.0,3.8,BLOB[2]); blob(s,12.7,0.4,2.6,BLOB[3],35)
    card(s,2.7,2.6,7.93,2.3,fill=WHITE,rad=0.3)
    text(s,2.8,2.6,7.73,1.5,[[('🎉 ',48,HOTPINK,True,BF),('Thank you!',52,HOTPINK,True,TF),(' '+EM[1],48,HOTPINK,True,BF)]],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
    text(s,2.8,3.95,7.73,0.7,[[(cfg['bye'],22,'7A5C9E',True,TF)]],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
    text(s,1.5,5.1,10.3,0.8,[[(' '.join(EM[3])+' ✨',40,INK,False,BF)]],align=PP_ALIGN.CENTER)
    logo(s)

    G.finish(prs,helpers)          # service screens go to the very end of the file
    def pos(sl): return prs.slides.index(sl)+1
    nums=[pos(sl) for sl in game_slides]
    games_txt=(f'{min(nums)}-{max(nums)}' if len(nums)>1 else (str(nums[0]) if nums else '-'))
    render_notes(s_notes,cfg,{'GAMES':games_txt,'VIDEO':str(pos(s_video)),'SPEAKING':str(pos(s_speak)),
                              'SERVICE':(f'{len(prs.slides._sldIdLst)-len(helpers)+1}-{len(prs.slides._sldIdLst)}' if helpers else '-')})
    prs.save(out)
    print('saved',os.path.basename(out),'|',len(prs.slides._sldIdLst),'slides |',
          (f'games {games_txt}, service screens at the end' if nums else 'no games'))
    return prs

def render_notes(s,cfg,tokens):
    """Teacher notes. {GAMES} / {VIDEO} / {SPEAKING} / {SERVICE} in any note text
    are replaced with the real slide numbers, so the notes can never drift."""
    def sub(t):
        for k,v in tokens.items(): t=t.replace('{'+k+'}',v)
        return t
    blob(s,0.3,0.4,2.6,'5A4A7A',50); blob(s,13.0,7.2,3.0,'5A4A7A',50)
    text(s,0.7,0.35,11.9,0.9,[[('🔒 ',30,GOLD,True,BF),('ONLY FOR THE TEACHER!!!!! ',30,GOLD,True,TF),('👩\u200d🏫',30,GOLD,True,BF)]],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
    card(s,0.7,1.35,11.93,5.4,fill='3A2E55',rad=0.05,sh=False)
    tb=s.shapes.add_textbox(Inches(1.0),Inches(1.5),Inches(11.3),Inches(5.1)); tf=tb.text_frame
    tf.word_wrap=True; tf.vertical_anchor=MSO_ANCHOR.TOP
    tf.margin_left=Pt(4); tf.margin_right=Pt(4); tf.margin_top=Pt(4)
    notes=list(cfg['notes'])
    if cfg.get('games') and not any('{GAMES}' in n.get('label','')+n.get('text','') for n in notes):
        notes.append({'label':f'SLIDES {tokens["GAMES"]}:','text':GAME_NOTE})
    first=True
    for nt in notes:
        p=tf.paragraphs[0] if first else tf.add_paragraph(); first=False
        p.space_after=Pt(5); p.line_spacing=1.0
        r=p.add_run(); r.text=sub(nt['label'])+' '; r.font.size=Pt(11.5); r.font.bold=True; r.font.name=BF; r.font.color.rgb=C(GOLD)
        r=p.add_run(); r.text=sub(nt['text']); r.font.size=Pt(11.5); r.font.name=BF; r.font.color.rgb=C('F2ECFF')
        if nt.get('url'):
            p.add_run().text=' '; lnk(p,nt['url'],nt['url'],11.5,'9EC9FF',BF)
        if nt.get('tail'):
            r=p.add_run(); r.text=sub(nt['tail']); r.font.size=Pt(11.5); r.font.name=BF; r.font.color.rgb=C('F2ECFF')
        if nt.get('red'):
            r=p.add_run(); r.text='  🟥 '+sub(nt['red']); r.font.size=Pt(11.5); r.font.bold=True; r.font.name=BF; r.font.color.rgb=C('FF6B6B')
    logo(s,dark=True)

def text_in_shape(sh,txt,size,col,bold=True,font=TF):
    """Put the text INSIDE the shape - never a separate textbox on top of it
    (a textbox on top swallows the click and kills the centre of the button)."""
    tf=sh.text_frame; tf.word_wrap=True; tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    for m in ('left','right','top','bottom'): setattr(tf,f'margin_{m}',Pt(3))
    p=tf.paragraphs[0]; p.alignment=PP_ALIGN.CENTER
    r=p.add_run(); r.text=txt; r.font.size=Pt(size); r.font.bold=bold; r.font.name=font; r.font.color.rgb=C(col)
    return sh
