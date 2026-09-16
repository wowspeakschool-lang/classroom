#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
WOW Speak in-deck games. No WordWall, no internet, no accounts.

Two mechanisms underneath:
  * slide-to-slide hyperlinks (work in every viewer);
  * real PowerPoint "Appear / Disappear on click" animations, written straight
    into the slide's <p:timing>.

THE RULES (each one was paid for - do not break them)
 1. One card = one shape, text INSIDE the shape (btn / card_text). A label on
    top of a button eats the click (only the thin rim stays alive) and an
    animated card made of three objects needs three clicks to vanish.
    A picture on top is allowed, but it needs the same link, and for animation
    the card and the picture are passed as ONE group.
 2. The task must be honest - captions must not give the answer away. Captions
    live only on the New Words slide. No captions and no give-away emoji in the
    quiz, in "what's missing" and on the board.
 3. One round PER WORD, not one round per lesson. Right -> next word, wrong ->
    same question again. Counter 1/6 ... 6/6 in the corner, final slide at the end.
 4. The correct answer never walks 1-2-3-4: its position is (k*3) % nopt.
 5. A click past a button must not flip the slide -> no_click_advance() on the
    quiz and on the board. The teacher moves on with the arrow key.
 6. Service screens go to the END of the file via finish(prs, helpers).
 7. Feedback screens come back on a click ANYWHERE (link_whole_slide).
 8. A card must never be the colour of the slide background -> pal(bg, i),
    never PASTELS[i % 6].
 9. Rewrite the teacher notes for the games (F5 only, click past a button does
    not advance, captions removed on purpose).
10. The deck gets longer. If the lesson does not fit, cut rounds knowingly
    (new or hard words only).
"""
import os
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.oxml import parse_xml

import wowspeak_builder as B

PICS=('.png','.jpg','.jpeg','.webp','.gif','.bmp')
# Card colours. Deliberately NOT B.PASTELS - those are the slide backgrounds.
CARD_COLOURS=['FFB6D9','A9D6FF','FFE89A','C9B3F0','9EE9C6','FFC79E']
GREEN='2E9E5B'; AMBER='E8A33D'

# ----------------------------------------------------------------- basics
def pal(bgc,i):
    """Card colour #i that can never come out equal to the slide background."""
    bgc=(bgc or '').upper().lstrip('#')
    cand=[c for c in CARD_COLOURS if c.upper()!=bgc] or CARD_COLOURS
    return cand[i%len(cand)]

def card_text(sh,txt,size=20,col=None,bold=True,font=None,emoji=''):
    """Write the text INSIDE the shape (rule 1)."""
    col=col or B.INK
    tf=sh.text_frame; tf.word_wrap=True; tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    for m in ('left','right','top','bottom'): setattr(tf,f'margin_{m}',Pt(3))
    p=tf.paragraphs[0]; p.alignment=PP_ALIGN.CENTER
    if emoji:
        r=p.add_run(); r.text=emoji+' '; r.font.size=Pt(size); r.font.bold=bold; r.font.name=B.BF; r.font.color.rgb=B.C(col)
    r=p.add_run(); r.text=txt; r.font.size=Pt(size); r.font.bold=bold
    r.font.name=font or (B.BF if _is_emoji_only(txt) else B.TF); r.font.color.rgb=B.C(col)
    return sh

def _is_emoji_only(t):
    return bool(t) and all(ord(c)>0x2190 or c==' ' for c in t)

def btn(slide,x,y,w,h,txt='',fill=None,size=20,col=None,target=None,url=None,rad=0.14,bold=True,font=None,shade=True):
    """One shape = one card = one button. Text lives inside it."""
    sh=B.card(slide,x,y,w,h,fill=fill or B.WHITE,rad=rad,sh=shade)
    if txt: card_text(sh,txt,size=size,col=col,bold=bold,font=font)
    if target is not None: sh.click_action.target_slide=target
    if url: sh.click_action.hyperlink.address=url
    return sh

def media_on(slide,sh,media,pad=0.22,size=60,target=None,url=None,panel=True):
    """Put a picture (or a big emoji) on the card. Everything laid on top of the
    card swallows the click, so every piece gets the SAME link (rule 1) and the
    whole lot is returned as one list - that is also the animation group.
    The picture sits on a white panel: icons are cut from a white sheet, and
    without the panel a bare white rectangle shows on the coloured card."""
    if not media: return []
    x,y=sh.left/914400.0,sh.top/914400.0
    w,h=sh.width/914400.0,sh.height/914400.0
    if isinstance(media,str) and media.lower().endswith(PICS):
        if not os.path.exists(media): return []
        out=[]
        if panel: out.append(B.card(slide,x+pad,y+pad,w-2*pad,h-2*pad,fill=B.WHITE,rad=0.12,sh=False))
        q=pad+(0.14 if panel else 0.0)
        out.append(B.fit(slide,media,x+q,y+q,w-2*q,h-2*q))
        if target is None and url is None:
            try: target=sh.click_action.target_slide
            except Exception: target=None
        for e in out:
            if target is not None: e.click_action.target_slide=target
            elif url: e.click_action.hyperlink.address=url
        return out
    card_text(sh,media,size=size,font=B.BF)   # emoji straight into the shape
    return []

def _bg_of(slide,default=B.WHITE):
    try: return str(slide.background.fill.fore_color.rgb)
    except Exception: return default

def link_whole_slide(slide,target,bg=None):
    """Whole slide becomes one big button (rule 7). The rectangle is pushed to
    the back and painted in the background colour, so it is invisible but
    clickable everywhere - a no-fill shape only reacts on its rim."""
    sh=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,B.Emu(12192000),B.Emu(6858000))
    sh.fill.solid(); sh.fill.fore_color.rgb=B.C(bg or _bg_of(slide))
    sh.line.fill.background(); sh.shadow.inherit=False
    sh.click_action.target_slide=target
    el=sh._element; tree=el.getparent(); tree.remove(el); tree.insert(2,el)  # to the back
    return sh

def no_click_advance(slide):
    """Rule 5: a click past a button must not flip the slide. Animations still
    run on click - only the slide advance is switched off. Teacher uses the arrow."""
    sld=slide._element
    for t in sld.findall(qn('p:transition')): sld.remove(t)
    tr=parse_xml('<p:transition xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" advClick="0"/>')
    tim=sld.find(qn('p:timing'))
    if tim is not None: tim.addprevious(tr)
    else: sld.append(tr)
    return tr

# ------------------------------------------------------- click animations
_P='http://schemas.openxmlformats.org/presentationml/2006/main'

def _eff(ids,spid,kind,first):
    cid=next(ids); bid=next(ids)
    cls='entr' if kind=='entr' else 'exit'
    val='visible' if kind=='entr' else 'hidden'
    node='clickEffect' if first else 'withEffect'
    return (f'<p:par><p:cTn id="{cid}" presetID="1" presetClass="{cls}" presetSubtype="0" '
            f'fill="hold" grpId="0" nodeType="{node}">'
            f'<p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst><p:set><p:cBhvr>'
            f'<p:cTn id="{bid}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>'
            f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>'
            f'<p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr>'
            f'<p:to><p:strVal val="{val}"/></p:to></p:set></p:childTnLst></p:cTn></p:par>')

def animate(slide,steps):
    """steps = [(kind, [shape|group of shapes]), ...] - one mouse click per step.
    kind: 'entr' (Appear) or 'exit' (Disappear). Everything inside one step
    happens on the same click (rule 1: card + picture travel together)."""
    steps=[s for s in steps if s[1]]
    if not steps: return None
    ids=iter(range(2,100000)); seq=next(ids); clicks=[]
    for kind,shapes in steps:
        if not isinstance(shapes,(list,tuple)): shapes=[shapes]
        spids=[sh.shape_id for sh in shapes if sh is not None]
        if not spids: continue
        outer=next(ids); inner=next(ids)
        effs=''.join(_eff(ids,sp,kind,i==0) for i,sp in enumerate(spids))
        clicks.append(f'<p:par><p:cTn id="{outer}" fill="hold">'
                      f'<p:stCondLst><p:cond delay="indefinite"/></p:stCondLst><p:childTnLst>'
                      f'<p:par><p:cTn id="{inner}" fill="hold">'
                      f'<p:stCondLst><p:cond delay="0"/></p:stCondLst>'
                      f'<p:childTnLst>{effs}</p:childTnLst></p:cTn></p:par>'
                      f'</p:childTnLst></p:cTn></p:par>')
    cond=('<p:cond evt="{e}" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond>')
    xml=(f'<p:timing xmlns:p="{_P}"><p:tnLst><p:par>'
         f'<p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
         f'<p:seq concurrent="1" nextAc="seek"><p:cTn id="{seq}" dur="indefinite" nodeType="mainSeq">'
         f'<p:childTnLst>{"".join(clicks)}</p:childTnLst></p:cTn>'
         f'<p:prevCondLst>{cond.format(e="onPrev")}</p:prevCondLst>'
         f'<p:nextCondLst>{cond.format(e="onNext")}</p:nextCondLst>'
         f'</p:seq></p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>')
    sld=slide._element
    for t in sld.findall(qn('p:timing')): sld.remove(t)
    tim=parse_xml(xml); sld.append(tim)
    return tim

# ------------------------------------------------------------- structure
def finish(prs,helpers):
    """Rule 6: push every service screen to the very end of the file, so that
    arrowing through the lesson never lands on 'Correct / Try again'."""
    lst=prs.slides._sldIdLst; els=list(lst)
    by_id={prs.slides[i].slide_id:els[i] for i in range(len(els))}
    seen=set()
    for sl in helpers:
        if sl is None or sl.slide_id in seen: continue
        seen.add(sl.slide_id); el=by_id[sl.slide_id]
        lst.remove(el); lst.append(el)
    return prs

def _head(slide,title,bg,counter=None,dark=False):
    B.band(slide,title,dark=dark,size=28)
    if counter:
        c=B.card(slide,11.75,0.5,1.35,0.7,fill=B.WHITE if not dark else '3A2E55',rad=0.4)
        card_text(c,counter,size=19,col=B.HOTPINK if not dark else B.GOLD)
    B.logo(slide,dark=dark)

# ------------------------------------------------------------ THE GAMES
def g_quiz_series(prs,bg,items,helpers=None,title=None,mode='picture',nopt=4,
                  final='🏆 Well done! You know all the words!'):
    """A round for EVERY word (rule 3).
    items : [(word, media)] for mode='picture' (media = icon path or emoji)
            [(word, definition)] for mode='word'
    Right answer -> 'Correct' -> next word. Wrong -> 'Try again' -> same word."""
    helpers=[] if helpers is None else helpers
    n=len(items); nopt=max(2,min(nopt,n))
    qs=[B.new_slide(prs,bg) for _ in range(n)]
    fin=B.new_slide(prs,bg)
    outs=[]
    for k in range(n):
        ok=B.new_slide(prs,GREEN); bad=B.new_slide(prs,AMBER)
        outs.append((ok,bad)); helpers+=[ok,bad]
    for k,(word,extra) in enumerate(items):
        s=qs[k]; ok,bad=outs[k]
        nxt=qs[k+1] if k+1<n else fin
        q=f'Which one is "{word}"?' if mode=='picture' else f'Which word means: "{extra}"?'
        _head(s,title or ('🎯 Point & Race!' if mode=='picture' else '🎯 Word hunt!'),bg,counter=f'{k+1}/{n}')
        B.text(s,0.9,1.62,11.5,0.7,[[(q,22,'5A4A7A',True,B.TF)]],align=PP_ALIGN.CENTER)
        pos=(k*3)%nopt                                  # rule 4
        order=[None]*nopt; order[pos]=k
        others=[(k+1+j)%n for j in range(n-1)][:nopt-1]
        for slot in range(nopt):
            if order[slot] is None: order[slot]=others.pop(0)
        cw=min(2.75,(11.6-(nopt-1)*0.35)/nopt); gap=0.35
        gx=(B.SLIDE_W-(nopt*cw+(nopt-1)*gap))/2
        gy,ch=(2.5,3.3) if mode=='picture' else (3.0,2.1)
        for slot,idx in enumerate(order):
            x=gx+slot*(cw+gap)
            tgt=ok if idx==k else bad
            c=btn(s,x,gy,cw,ch,fill=pal(bg,slot),target=tgt)
            if mode=='picture': media_on(s,c,items[idx][1],target=tgt)
            else: card_text(c,items[idx][0],size=24,col=B.INK)
        no_click_advance(s)                              # rule 5
        # feedback screens - back on a click anywhere (rule 7)
        link_whole_slide(ok,nxt,GREEN)
        B.text(ok,1.0,2.5,11.3,1.4,[[('YES! Correct! 🎉',54,B.GOLD,True,B.TF)]],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
        B.text(ok,1.0,4.1,11.3,0.8,[[('Click anywhere to go on ➡️',20,B.WHITE,True,B.BF)]],align=PP_ALIGN.CENTER)
        link_whole_slide(bad,s,AMBER)
        B.text(bad,1.0,2.5,11.3,1.4,[[('Try again! 🤔',54,B.WHITE,True,B.TF)]],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
        B.text(bad,1.0,4.1,11.3,0.8,[[('Click anywhere to come back ↩️',20,'FFF6E0',True,B.BF)]],align=PP_ALIGN.CENTER)
    B.blob(fin,0.4,0.6,3.2,B.BLOB[2]); B.blob(fin,12.8,7.0,3.6,B.BLOB[1])
    c=B.card(fin,2.6,2.3,8.13,2.4,fill=B.WHITE,rad=0.25)
    card_text(c,final,size=30,col=B.HOTPINK)
    B.text(fin,1.5,5.1,10.3,0.8,[[('⭐ 🎉 ⭐ 🎉 ⭐',40,B.INK,False,B.BF)]],align=PP_ALIGN.CENTER)
    B.logo(fin)
    return helpers

def g_quiz(prs,bg,items,helpers=None,**kw):
    """Same quiz on emoji instead of icons (items = [(word, '🐶')])."""
    return g_quiz_series(prs,bg,items,helpers,**kw)

def g_missing_pics(prs,bg,items,helpers=None,title='👀 What is missing?',order=None,rounds=None):
    """The card really disappears (PowerPoint 'Disappear on click').
    items = [(word, media)]. One click = one card gone; the class names it."""
    helpers=[] if helpers is None else helpers
    s=B.new_slide(prs,bg)
    _head(s,title,bg)
    B.text(s,0.9,1.6,11.5,0.6,[[('Look and remember. Then say what disappeared!',18,'5A4A7A',True,B.TF)]],align=PP_ALIGN.CENTER)
    n=len(items); cols=3 if n>4 else n
    cw=min(3.3,(11.5-(cols-1)*0.4)/cols); gap=0.4; rows=(n+cols-1)//cols
    ch=min(2.2,(4.45-(rows-1)*0.3)/rows)          # keep clear of the logo line
    gx=(B.SLIDE_W-(cols*cw+(cols-1)*gap))/2; gy=2.35
    groups=[]
    for i,(word,media) in enumerate(items):
        r=i//cols; c=i%cols; x=gx+c*(cw+gap); y=gy+r*(ch+0.3)
        sh=B.card(s,x,y,cw,ch,fill=pal(bg,i))        # rule 8
        groups.append([sh]+media_on(s,sh,media,size=66))        # rule 1: one click
    seq=order if order is not None else [(i*3+1)%n for i in range(n)]
    seen=[]; [seen.append(i) for i in seq if i not in seen and 0<=i<n]
    seen+=[i for i in range(n) if i not in seen]
    if rounds: seen=seen[:rounds]
    animate(s,[('exit',groups[i]) for i in seen])
    B.logo(s)
    return helpers

def g_missing_anim(prs,bg,items,helpers=None,**kw):
    """Same game on emoji (items = [(word,'🍂')]) - no icons needed."""
    return g_missing_pics(prs,bg,items,helpers,**kw)

def g_reveal(prs,bg,lines,helpers=None,title='💡 Answers',subtitle=None):
    """Answers appear one per click. lines = ['...', ...] or [(left, right)]."""
    helpers=[] if helpers is None else helpers
    s=B.new_slide(prs,bg)
    _head(s,title,bg)
    if subtitle:
        B.text(s,0.9,1.6,11.5,0.6,[[(subtitle,18,'5A4A7A',True,B.TF)]],align=PP_ALIGN.CENTER)
    n=len(lines); top=2.3; h=min(0.95,(4.9-(n-1)*0.18)/max(n,1)); gap=0.18
    steps=[]
    for i,ln in enumerate(lines):
        y=top+i*(h+gap)
        if isinstance(ln,(list,tuple)):
            left,right=ln
            B.text(s,1.3,y,5.2,h,[[(left,19,B.INK,True,B.BF)]],align=PP_ALIGN.LEFT,anchor=MSO_ANCHOR.MIDDLE)
            sh=B.card(s,6.6,y,5.4,h,fill=pal(bg,i)); card_text(sh,right,size=19,col=B.INK)
        else:
            sh=B.card(s,2.2,y,8.9,h,fill=pal(bg,i)); card_text(sh,ln,size=19,col=B.INK)
        steps.append(('entr',[sh]))
    animate(s,steps)
    B.logo(s)
    return helpers

def g_board(prs,bg,questions,helpers=None,title='🎲 Choose a number!',back='⬅️ Back to the board'):
    """Clickable board: number -> question -> back. Question screens are service
    slides and end up at the end of the file (rule 6)."""
    helpers=[] if helpers is None else helpers
    board=B.new_slide(prs,bg)
    _head(board,title,bg)
    n=len(questions); rows=max(1,-(-n//5)); cols=-(-n//rows)
    cw=min(2.4,(11.4-(cols-1)*0.4)/cols); ch=min(2.0,(4.7-(rows-1)*0.35)/rows)
    gx=(B.SLIDE_W-(cols*cw+(cols-1)*0.4))/2
    gy=2.0+max(0.0,(4.7-(rows*ch+(rows-1)*0.35))/2)
    for i,q in enumerate(questions):
        qs=B.new_slide(prs,bg); helpers.append(qs)
        r=i//cols; c=i%cols
        btn(board,gx+c*(cw+0.4),gy+r*(ch+0.35),cw,ch,txt=str(i+1),fill=pal(bg,i),size=40,target=qs)
        _head(qs,f'Question {i+1}',bg)
        card=B.card(qs,1.3,2.0,10.73,3.0,fill=B.WHITE,rad=0.1)
        card_text(card,q,size=26,col=B.INK,bold=True)
        btn(qs,4.9,5.4,3.53,0.85,txt=back,fill=pal(bg,i),size=17,target=board)
        no_click_advance(qs)                              # rule 5
    no_click_advance(board)
    return helpers

# --------------------------------------------------- builder entry point
def play(prs,spec,cfg):
    """cfg['games'] = [{'type':'quiz'}, {'type':'missing'}, ...] -> slides."""
    t=spec.get('type','quiz'); bg=spec.get('bg') or B.PASTELS['yellow']
    lvl=cfg.get('level','beg'); vocab=cfg.get('vocab',[])
    items=spec.get('items') or vocab
    mode=spec.get('mode') or ('word' if lvl=='adv' else 'picture')
    def kw(*names): return {k:v for k,v in spec.items() if k in names}
    if t=='quiz':
        return g_quiz_series(prs,bg,items,mode=mode,**kw('title','nopt','final'))
    if t=='missing':
        if mode=='word':
            emo=spec.get('emoji',{}); items=[(w,emo.get(w,'💬')) for w,_ in items]
        return g_missing_pics(prs,bg,items,**kw('title','order','rounds'))
    if t=='reveal':
        return g_reveal(prs,bg,spec.get('lines') or [w for w,_ in items],**kw('title','subtitle'))
    if t=='board':
        return g_board(prs,bg,spec.get('questions',[]),**kw('title','back'))
    raise ValueError(f'unknown game type: {t}')
