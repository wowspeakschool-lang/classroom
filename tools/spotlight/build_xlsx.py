"""tsv/grade_N.tsv → docs/Spotlight_словари_2-11.xlsx: сводка + лист на класс."""
import csv, collections
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
wb = Workbook(); ov = wb.active; ov.title = 'Сводка'
ov.append(['Класс', 'Модуль', 'Раздел (юнит)', 'Слов', 'с переводом'])
hdr = ['Класс', 'Модуль', 'Раздел (юнит)', 'Слово', 'Часть речи', 'Перевод', 'Примечание']
bold = Font(bold=True); fill = PatternFill('solid', fgColor='DDEBF7')
tot = 0
for g in range(2, 12):
    rows = list(csv.reader(open(f'tsv/grade_{g}.tsv'), delimiter='\t'))[1:]
    ws = wb.create_sheet(f'{g} класс'); ws.append(hdr)
    for r in rows: ws.append([int(r[0])] + r[1:])
    for c, w in zip('ABCDEFG', (7, 26, 34, 34, 11, 50, 26)): ws.column_dimensions[c].width = w
    for c in ws[1]: c.font = bold; c.fill = fill
    ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
    secs = collections.OrderedDict()
    for r in rows:
        a = secs.setdefault((r[1], r[2]), [0, 0]); a[0] += 1; a[1] += bool(r[5].strip())
    for (m, s), (n, t) in secs.items(): ov.append([g, m, s, n, t])
    tot += len(rows); print(g, len(rows))
for c in ov[1]: c.font = bold; c.fill = fill
for c, w in zip('ABCDE', (7, 40, 44, 8, 12)): ov.column_dimensions[c].width = w
ov.freeze_panes = 'A2'; ov.auto_filter.ref = ov.dimensions
wb.save('/home/user/classroom/docs/Spotlight_словари_2-11.xlsx'); print('total', tot)
