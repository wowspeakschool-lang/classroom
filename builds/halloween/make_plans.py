"""Планы уроков Halloween в Word и PDF → lessons/halloween/plans/ (кнопки «Скачать план» в окне 📋).
Берёт PLAN_HTML из готовых уроков lessons/halloween/*.html. Нужны: pip install pypandoc_binary python-docx;
playwright (Chromium) для PDF. Запуск: python3 builds/halloween/make_plans.py"""
import re, json, os, subprocess, glob, pypandoc, docx
from docx.enum.section import WD_ORIENT
from docx.shared import Cm, Pt, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
REPO = os.path.abspath(os.path.join(os.path.dirname(__file__), '../..'))
L, OUT, W = REPO + '/lessons/halloween/', REPO + '/lessons/halloween/plans/', '/tmp/halloween_plans/'
os.makedirs(OUT, exist_ok=True); os.makedirs(W, exist_ok=True)
P = pypandoc.get_pandoc_path()
subprocess.run([P, '-o', W + 'ref.docx', '--print-default-data-file', 'reference.docx'], check=True)
d = docx.Document(W + 'ref.docx')
for s in d.sections:
    s.orientation = WD_ORIENT.LANDSCAPE; s.page_width, s.page_height = Cm(29.7), Cm(21.0)
    s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Cm(1.5)
for n in ['Normal', 'Body Text', 'First Paragraph', 'Compact']:
    try: d.styles[n].font.name = 'Arial'; d.styles[n].font.size = Pt(10)
    except KeyError: pass
d.save(W + 'ref_land.docx')
for f in sorted(glob.glob(L + '*.html')):
    name = 'halloween_' + os.path.basename(f)[:-5]
    h = json.loads(re.search(r'const PLAN_HTML=("(?:[^"\\]|\\.)*");', open(f).read()).group(1))
    h = re.sub(r'<p class="dl".*?</p>', '', h, flags=re.S)
    open(W + name + '.html', 'w').write(h)
    open(W + 'in.html', 'w').write(re.sub(r'<title>[^<]*</title>', '', h))
    subprocess.run([P, W + 'in.html', '-f', 'html', '-t', 'docx', '--reference-doc=' + W + 'ref_land.docx', '-o', OUT + name + '.docx'], check=True)
    dd = docx.Document(OUT + name + '.docx')
    for t in dd.tables:
        pr = t._tbl.tblPr; b = OxmlElement('w:tblBorders')
        for e in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
            x = OxmlElement('w:' + e); x.set(qn('w:val'), 'single'); x.set(qn('w:sz'), '4'); x.set(qn('w:color'), 'BFB3D9'); b.append(x)
        for o in pr.findall(qn('w:tblBorders')): pr.remove(o)
        pr.append(b)
        for c in t.rows[0].cells:
            sh = OxmlElement('w:shd'); sh.set(qn('w:val'), 'clear'); sh.set(qn('w:fill'), '6E3510'); c._tc.get_or_add_tcPr().append(sh)
            for p in c.paragraphs:
                for r in p.runs: r.font.bold = True; r.font.color.rgb = RGBColor(255, 255, 255)
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    for r in p.runs: r.font.size = Pt(9.5)
    dd.save(OUT + name + '.docx')
js = """const {chromium}=require('playwright');const fs=require('fs');(async()=>{const b=await chromium.launch();const p=await b.newPage();
for(const f of fs.readdirSync('%s').filter(x=>x.startsWith('halloween_')&&x.endsWith('.html'))){await p.goto('file://%s'+f);
await p.addStyleTag({content:'body{background:#fff!important;max-width:none!important;padding:0!important}table,.meta,.voc{box-shadow:none!important}tr{page-break-inside:avoid}'});
await p.pdf({path:'%s'+f.replace('.html','.pdf'),format:'A4',landscape:true,printBackground:true,margin:{top:'12mm',bottom:'12mm',left:'12mm',right:'12mm'}});}
await b.close();})();""" % (W, W, OUT)
open(W + 'pdf.js', 'w').write(js)
subprocess.run(['node', W + 'pdf.js'], check=True)
print(sorted(os.listdir(OUT)))
