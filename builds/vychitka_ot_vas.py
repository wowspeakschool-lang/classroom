# -*- coding: utf-8 -*-
"""Общий документ «Вычитка курсов — что нужно от вас» из markdown в Word.

    python3.13 builds/vychitka_ot_vas.py

Источник — `builds/vychitka_ot_vas.md`: по разделу на курс, каждый чат правит
только свой. Документ пересобирается после любой правки источника.

Разметка простая: # заголовок, ## раздел, **жирный**, `код`, списки «- »,
таблицы с «|», «---» — разделитель. Стиль тот же, что у чеклиста методиста.
"""
import io, re, sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

SRC = sys.argv[1] if len(sys.argv) > 1 else "builds/vychitka_ot_vas.md"
OUT = sys.argv[2] if len(sys.argv) > 2 else "builds/Вычитка — что нужно от вас.docx"
INK   = RGBColor(0x1F, 0x29, 0x37)
HEADC = RGBColor(0x14, 0x3A, 0x6B)
ACC   = RGBColor(0xB1, 0x3E, 0x0F)
MUTED = RGBColor(0x5A, 0x63, 0x72)
CODE  = RGBColor(0x0B, 0x61, 0x6E)
HEAD_FILL, SOFT_FILL, GRID = '1F3A6B', 'F3F7FC', 'C9D4E4'

doc = Document()
s = doc.sections[0]
s.left_margin = s.right_margin = Cm(1.8)
s.top_margin = s.bottom_margin = Cm(1.5)
doc.styles['Normal'].font.name = 'Calibri'
doc.styles['Normal'].font.size = Pt(10.5)
doc.styles['Normal'].font.color.rgb = INK

TOKEN = re.compile(r'(\*\*[^*]+\*\*|`[^`]+`)')


def runs(par, text, size=10.5, base=None, italic=False):
    for part in TOKEN.split(text):
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            r = par.add_run(part[2:-2]); r.bold = True
            r.font.color.rgb = ACC if part.startswith('**🔴') else (base or HEADC)
        elif part.startswith('`') and part.endswith('`'):
            r = par.add_run(part[1:-1]); r.font.name = 'Consolas'
            r.font.color.rgb = CODE
        else:
            r = par.add_run(part); r.font.color.rgb = base or INK
            r.italic = italic
        r.font.size = Pt(size)


def para(text='', *, size=10.5, after=6, base=None, bullet=False, italic=False):
    p = doc.add_paragraph(style='List Bullet' if bullet else None)
    p.paragraph_format.space_after = Pt(after)
    if text:
        runs(p, text, size, base, italic)
    return p


def shade(cell, fill):
    el = OxmlElement('w:shd'); el.set(qn('w:val'), 'clear')
    el.set(qn('w:color'), 'auto'); el.set(qn('w:fill'), fill)
    cell._tc.get_or_add_tcPr().append(el)


def table(rows):
    head, body = rows[0], rows[1:]
    t = doc.add_table(rows=len(rows), cols=len(head))
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), '4')
        el.set(qn('w:space'), '0'); el.set(qn('w:color'), GRID)
        borders.append(el)
    t._tbl.tblPr.append(borders)
    for ci, h in enumerate(head):
        c = t.rows[0].cells[ci]; shade(c, HEAD_FILL); c.text = ''
        p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h.strip('* ')); r.bold = True; r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    for ri, row in enumerate(body, start=1):
        for ci, val in enumerate(row):
            c = t.rows[ri].cells[ci]; c.text = ''
            if ci == 0:
                shade(c, SOFT_FILL)
            p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(1)
            runs(p, val, 9.5, HEADC if ci == 0 else INK)
    return t


import datetime
lines = io.open(SRC, encoding='utf-8').read().split('\n')
i, buf, tbl = 0, [], []


def flush():
    global buf
    if buf:
        para(' '.join(buf))
        buf = []


while i < len(lines):
    ln = lines[i].rstrip()
    if ln.startswith('|'):
        tbl.append([c.strip() for c in ln.strip('|').split('|')])
        i += 1
        continue
    if tbl:
        table([r for r in tbl if not all(set(c) <= set('-: ') for c in r)])
        para('', after=4)
        tbl = []
    if not ln.strip() or ln.strip() == '---':
        flush()
    elif ln.startswith('# '):
        flush(); para(ln[2:], size=19, after=4, base=HEADC)
    elif ln.startswith('### '):
        flush()
        p3 = para(ln[4:], size=11.5, after=4, base=ACC)
        p3.paragraph_format.space_before = Pt(10)
    elif ln.startswith('## '):
        flush()
        p = para(ln[3:], size=14, after=5, base=HEADC)
        p.paragraph_format.space_before = Pt(14)
    elif ln.startswith('- ') or ln.startswith('* '):
        flush(); para(ln[2:], bullet=True, after=3)
    elif ln.startswith('  ') and buf:
        buf.append(ln.strip())
    else:
        buf.append(ln.strip())
    i += 1
flush()
if tbl:
    table([r for r in tbl if not all(set(c) <= set('-: ') for c in r)])

para('Документ пересобран ' + datetime.date.today().strftime('%d.%m.%Y') +
     '. Источник: builds/vychitka_ot_vas.md', size=9, after=0, base=MUTED, italic=True)
doc.save(OUT)
print('документ сохранён:', OUT)
