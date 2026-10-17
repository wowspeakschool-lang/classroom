# -*- coding: utf-8 -*-
"""Документ «Что нужно от вас — SM0, юнит 0».

    python3.13 builds/sm0/u0/todo.py "builds/sm0/u0/Что нужно от вас — SM0.docx"

Только открытые пункты: что сделать руками и что решить. Сделанное не перечисляем.
"""
import sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT = sys.argv[1]
doc = Document()
for m in ('left_margin', 'right_margin', 'top_margin', 'bottom_margin'):
    setattr(doc.sections[0], m, Cm(1.8))
doc.styles['Normal'].font.name = 'Calibri'
doc.styles['Normal'].font.size = Pt(11)


def para(text, *, size=11, bold=False, italic=False, color=None, after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text)
    r.font.size = Pt(size); r.bold = bold; r.italic = italic
    if color:
        r.font.color.rgb = color
    return p


def table(headers, rows, widths):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = 'Table Grid'
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]; c.text = ''
        r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(10)
    for ri, row in enumerate(rows, start=1):
        for ci, v in enumerate(row):
            c = t.rows[ri].cells[ci]; c.text = ''
            for li, line in enumerate(str(v).split('\n')):
                p = c.paragraphs[0] if li == 0 else c.add_paragraph()
                p.paragraph_format.space_after = Pt(1)
                rr = p.add_run(line); rr.font.size = Pt(10)
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w)
    return t


para('Что нужно от вас — SM0, юнит 0', size=18, bold=True, after=2)
para('Файлы: план урока Lesson 6 (Word), презентация к нему (PowerPoint), '
     'ДЗ 1 и тест юнита (SQL). В базу ничего не залито — ждём переезда на WowSpeak Team.',
     size=10, italic=True, after=12)

para('Сделать руками', size=13, bold=True, after=4)
table(['№', 'Что', 'Где', 'Без этого'],
      [['1', 'Записать звук /s/ и /æ/ — звук, не имя буквы',
        'ДЗ 1, блок 2, вопросы 1 и 2 (поле audio пустое)',
        'Задание не работает: ребёнку нечего слушать'],
       ['2', 'Нажать «Озвучить пачкой»',
        'ДЗ 1 блок 4 и тест блок 6 — фразы Hello! / What’s your name? / '
        'How are you? / I’m Misty помечены audio_tts',
        'Вопросы без звука'],
       ['3', 'Открыть презентацию Lesson 6 в PowerPoint и прокликать игры',
        'Слайд 5 — шарики, слайд 18 — карточки',
        'Не увидим, если переход ведёт не туда'],
       ['4', 'Положить план и презентацию туда, где преподаватели берут планы',
        'Unit 0 Lesson 6.docx и Unit 0 Lesson 6.pptx',
        'Учителю нечего открыть на уроке'],
       ['5', 'Включить публикацию ДЗ 1 и теста — после того, как посмотрите их глазами',
        'На новом месте, после заливки',
        'Уроки не видны ученикам'],
       ['6', 'Вставить ссылку на тест в презентацию',
        'Слайд 25, плашка «paste the test link here»',
        'Учитель не запустит тест с урока']],
      [1.0, 5.6, 5.6, 4.8])

para('')
para('Решить', size=13, bold=True, after=4)
table(['№', 'Вопрос', 'Пояснение'],
      [['1', 'Менять ли старые картинки цветов на новые шарики',
        'В тесте и в домашках 2–5 цвета берутся из картинок выгрузки 92×150 — '
        'они мылятся. Новые шарики готовы (media/sm0/u0). Договорились смотреть '
        'это при вычитке SM0 — здесь просто чтобы не потерялось'],
       ['2', 'ДЗ 2 юнита 0, блок 7',
        '«Послушай цвет и выбери его», а варианты — русские слова. Это задание '
        'на перевод, которых мы не делаем. Правим при вычитке SM0'],
       ['3', 'Качество картинок всего курса',
        'В SM0 128 картинок из 150 меньше 300 px по большей стороне. '
        'При вычитке посчитаем, что перерисовываем, и соберём промпты листами']],
      [1.0, 5.0, 11.0])

doc.save(OUT)
print('документ сохранён:', OUT)
