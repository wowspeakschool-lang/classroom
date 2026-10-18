# -*- coding: utf-8 -*-
"""Документ «Что нужно от вас — GG3, юнит 0».

    python3.13 builds/gg3/u0/todo.py "builds/gg3/u0/Что нужно от вас — GG3.docx"

Только открытые пункты: что сделать руками и что решить. Сделанное не перечисляем.
"""
import sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT = sys.argv[1]
doc = Document()
for m in ('left_margin', 'right_margin', 'top_margin', 'bottom_margin'):
    setattr(doc.sections[0], m, Cm(1.8))
doc.styles['Normal'].font.name = 'Calibri'
doc.styles['Normal'].font.size = Pt(11)


def para(text, *, size=11, bold=False, italic=False, color=None, after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text)
    r.font.size = Pt(size); r.bold = bold; r.italic = italic
    if color:
        r.font.color.rgb = color
    return p


def table(headers, rows, widths):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = 'Table Grid'
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]; c.text = ''
        r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(10)
    for ri, row in enumerate(rows, start=1):
        for ci, v in enumerate(row):
            c = t.rows[ri].cells[ci]; c.text = ''
            for li, line in enumerate(str(v).split('\n')):
                p = c.paragraphs[0] if li == 0 else c.add_paragraph()
                p.paragraph_format.space_after = Pt(1)
                rr = p.add_run(line); rr.font.size = Pt(10)
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w)
    return t


para('Что нужно от вас — GG3, юнит 0', size=18, bold=True, after=2)
para('Файл один: тест юнита 0 (SQL, 21 блок). Домашек и планов в GG3 нет по '
     'вашему решению. В базу ничего не залито — ждём переезда на WowSpeak Team.',
     size=10, italic=True, after=12)

para('Сделать руками', size=13, bold=True, after=4)
table(['№', 'Что', 'Где', 'Без этого'],
      [['1', 'Нажать «Озвучить пачкой»',
        'Тест: предложения для «составь предложение» и диктант — всё помечено '
        'audio_tts. Блок аудирования кнопка не тронет: там audio_tts нет '
        'специально',
        'Задания на слух не работают'],
       ['2', 'Записать аудирование своим голосом',
        'Тест, блок 17 (LISTENING). Текст для записи — ниже в этом файле. '
        'Файл назвать gg3_u0_test_b17.mp3 и прислать мне — вставлю в блок',
        'Ученику нечего слушать, пять вопросов после рассказа не решаются'],
       ['3', 'Посмотреть глазами тест после заливки',
        'На новом месте',
        'Ошибку найдёт ученик, а не мы'],
       ['4', 'Включить публикацию',
        'Урок «Unit 0 Test», после проверки',
        'Тест не виден ученикам']],
      [1.0, 5.6, 5.6, 4.8])

para('')
para('Текст аудирования — для записи', size=13, bold=True, after=4)
para('Тест, блок 17. Читать спокойно, один раз; ученик слушает запись дважды '
     'сам. На экране текста нет.', size=10, italic=True, after=6)
para('This is Emma. Emma is thirteen years old. She lives with her mum and her '
     'gran Bertie. Emma usually goes to school by bike, but when it rains her '
     'mum drives her. Emma is good at dancing and she has got dance classes '
     'twice a week. This year she is going to study Geography more. Emma is '
     'scared of spiders. Her favourite food is pancakes: her gran always makes '
     'pancakes for breakfast on Sunday.', size=11, after=10)

para('')
para('Картинок в тесте нет: материал страниц 4–9 целиком текстовый. Состав — '
     'Present Simple и наречия частотности, город и транспорт, профессии и '
     'вопросительные слова, be going to и школа, прилагательные с предлогами, '
     'еда с some/any и контейнеры, плюс reading, listening и диктант.',
     size=9, italic=True)

doc.save(OUT)
print('документ сохранён:', OUT)
