# -*- coding: utf-8 -*-
"""Опись вебинаров с Google Диска («Вебинары» / 7-9, 10+): Вебинары_опись.xlsx из data/webinars.json.

    python3 spotlight/webinars.py
"""
import json, os, re, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from build import D, header, body

OUT = os.path.join(D, 'Вебинары_опись.xlsx')
FILL = {'7-9': PatternFill('solid', fgColor='FCE4D6'), '10+': PatternFill('solid', fgColor='DDEBF7')}
VID = {'есть': PatternFill('solid', fgColor='C6EFCE'), 'нет': PatternFill('solid', fgColor='FFC7CE')}
KIND = {'грамматика': PatternFill('solid', fgColor='E2EFDA'), 'лексика': PatternFill('solid', fgColor='FFF2CC')}


def role(name, mime):
    n = name.lower()
    ext = name.rsplit('.', 1)[-1].lower() if '.' in name else ''
    if 'folder' in mime:
        return 'вложенная папка'
    if 'video' in mime or ext in ('mp4', 'mov', 'avi', 'mkv', 'webm'):
        return 'видео'
    if 'audio' in mime or ext in ('mp3', 'wav', 'm4a'):
        return 'аудио'
    if 'selfstudy' in n.replace('_', '').replace('-', '').replace(' ', ''):
        what = 'самостоятельная'
    elif re.search(r'exercis|упражн', n):
        what = 'упражнения'
    elif re.search(r'(^|[_\W])(hw|homework)([_\W]|$)', n):
        what = 'домашнее задание'
    elif ext == 'docx' or ext == 'doc':
        what = 'план урока'
    elif re.search(r'lesson|presentation', n) or ext == 'pptx':
        what = 'презентация урока'
    else:
        what = 'файл'
    return f'{what} ({ext})' if ext else what


def main():
    items = json.load(open(os.path.join(D, 'data', 'webinars.json')))
    vmap = {}
    p = os.path.join(D, 'data', 'video_map.json')
    if os.path.exists(p):
        for v in json.load(open(p)):
            vmap.setdefault((v['group'], v['folder']), []).append(v)
    lists = {g: os.path.exists(os.path.join(D, 'data', f'videos_{g}.txt')) for g in ('7-9', '10+')}
    wb = Workbook()
    ws = wb.active; ws.title = 'Как читать'
    for line in [
        'Опись вебинаров (Google Диск → «Вебинары»)',
        '',
        'Строка = папка вебинара. «Название» — точное имя папки: так же будут называться видео, по нему потом сопоставляем.',
        '«Что разбирается» — по материалам внутри папки (план урока, презентация, самостоятельная): объём темы — формы, лица, значения, чего нет; замечания (дубли, ошибки, файлы не на месте).',
        '«Видео» / «Название видео» — есть ли запись вебинара (по списку видео от методиста; в самих папках видеофайлов нет). Сопоставление названий видео и папок — data/video_map.json; спорные места — в «Замечание к видео». «нет данных» — список видео этой группы ещё не прислан.',
        'Листы: «7-9», «10+», «Все» (с фильтрами) и «Видео без папки» — записи, к которым не нашлось папки с материалами.',
    ]:
        ws.append([line])
    ws['A1'].font = Font(bold=True, size=14)
    ws.column_dimensions['A'].width = 130
    for row in ws.iter_rows(min_row=2):
        row[0].alignment = Alignment(wrap_text=True)

    cols = ['№', 'Группа', 'Название (как папка)', 'Видео', 'Название видео', 'Тип', 'Тема (англ.)', 'Что разбирается', 'Материалы в папке', 'Замечание к видео', 'Ссылка на папку']
    W = [5, 7, 36, 9, 34, 11, 26, 90, 34, 30, 12]
    def sheet(title, rs):
        w = wb.create_sheet(title)
        header(w, cols, W)
        w.freeze_panes = 'D2'
        for i, e in enumerate(sorted(rs, key=lambda e: (e['group'] != '7-9', e['path'].lower())), 1):
            mats = '\n'.join(f"{role(f['name'], f.get('mimeType', ''))}: {f['name']}" for f in e['files'])
            url = f"https://drive.google.com/drive/folders/{e['folder_id']}"
            name = e['path'].split('/', 1)[1] if '/' in e['path'] else e['folder']
            vs = vmap.get((e['group'], e['folder']), [])
            have = 'есть' if vs or e['has_video'] else ('нет' if lists[e['group']] else 'нет данных')
            w.append([i, e['group'], name, have, '\n'.join(v['video'] for v in vs), e['kind'], e['topic'], e['comment'], mats,
                      '\n'.join(v['note'] for v in vs if v['note']), 'открыть'])
            c = w.cell(row=w.max_row, column=11)
            c.hyperlink = url
            c.font = Font(color='0563C1', underline='single')
        body(w)
        for row in w.iter_rows(min_row=2):
            row[1].fill = FILL.get(row[1].value, PatternFill())
            row[2].font = Font(bold=True)
            row[5].fill = KIND.get(row[5].value, PatternFill())
            row[3].fill = VID.get(row[3].value, PatternFill())
        w.auto_filter.ref = w.dimensions
    for g in ('7-9', '10+'):
        sheet(g, [e for e in items if e['group'] == g])
    sheet('Все', items)
    w = wb.create_sheet('Видео без папки')
    header(w, ['Группа', 'Название видео', 'Замечание'], [8, 50, 70])
    folders = {(e['group'], e['folder']) for e in items}
    p = os.path.join(D, 'data', 'video_map.json')
    for v in (json.load(open(p)) if os.path.exists(p) else []):
        if (v['group'], v['folder']) not in folders:
            w.append([v['group'], v['video'], v['note']])
    body(w)
    wb.save(OUT)
    print(OUT, {g: sum(e['group'] == g for e in items) for g in ('7-9', '10+')})


if __name__ == '__main__':
    main()
