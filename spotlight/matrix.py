# -*- coding: utf-8 -*-
"""Планировщик уроков: Spotlight_темы_уроков.xlsx.

Строка — подтема (урок), колонки — классы: модуль и статус (новое / расширение / повторение),
и отметки о видео отдельно для двух возрастных групп вебинаров.
Подтемы грамматики — data/subtopics/*.json; Use of English — из data/gN.json.
Отметки о видео — data/videos.json: {"<подтема>": {"7–9": "ссылка или название", "10+": "..."}}.

    python3 spotlight/matrix.py
"""
import json, os, re, sys, glob
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.worksheet.datavalidation import DataValidation
from build import D, CATS, LEX, load_rows, sort_rows, modnum, header, body, BOX, WRAP, ALIAS, CATFIX

OUT = os.path.join(D, 'Spotlight_темы_уроков.xlsx')

# 7–9 лет — Spotlight 2–4 (4 класс — по решению методиста: его темы идут в вебинарах 7–9). 10+ — с 5 класса.
JUNIOR = {2, 3, 4}
GROUPS = ['7–9', '10+']
def group(g):
    return '7–9' if g in JUNIOR else '10+'

FILL = {'7–9': PatternFill('solid', fgColor='FCE4D6'), '10+': PatternFill('solid', fgColor='DDEBF7')}
GREY = PatternFill('solid', fgColor='EDEDED')
YES = PatternFill('solid', fgColor='92D050')
NO = PatternFill('solid', fgColor='FF7C80')
SOON = PatternFill('solid', fgColor='FFC000')
STFILL = {'видео': YES, 'материалы': SOON, 'нет': NO}
STRANK = {'нет': 0, 'материалы': 1, 'видео': 2}

WF = [  # модели словообразования, когда в строке нет явных суффиксов
    ('national', 'national'), ('compound noun', 'compound nouns'), ('сложные существ', 'compound nouns'),
    ('compound adj', 'compound adjectives'), ('сложные прилаг', 'compound adjectives'),
    ('abstract', 'abstract nouns'), ('абстрактн', 'abstract nouns'),
    ('negative adj', 'negative adjectives'), ('отрицательных прилаг', 'negative adjectives'),
    ('nouns from adjectives', 'nouns from adjectives'), ('nouns from verbs', 'nouns from verbs'),
    ('verbs from', 'verbs from nouns/adjectives'), ('with prefixes', 'verbs with prefixes'),
    ('лиц', 'personal nouns'), ('приставки (toc', 'prefixes'), ('конверсия', 'conversion (a dress – to dress)'),
    ('revision', 'revision'), ('образование прилагательных (toc', 'forming adjectives'),
    ('образование глаголов (toc', 'forming verbs'),
]
AFFIX = re.compile(r'(?<![\w-])(-[a-z]+|[a-z]{1,3}-)(?![\w-])')


def lex_topic(r):
    s, t = r['scope'], r['topic']
    if t == 'Phrasal verbs':
        m = re.match(r'(\w+):', s) or re.search(r' с (\w+)', s) or re.match(r'(\w+) \+', s)
        return 'Use of English', 'Фразовые глаголы', f'Phrasal verbs: {m.group(1)}' if m else 'Phrasal verbs'
    if t == 'Word formation':
        low = s.lower()
        if 'national' in low or '-ese' in low:
            name = 'nationalities -an/-ish/-ian/-ese'
        elif 'participl' in low or ('-ed' in low and '-ing' in low and 'able' not in low):
            name = 'participle adjectives -ed/-ing'
        else:
            aff = [a for a in dict.fromkeys(AFFIX.findall(s)) if a not in ('-ing',) or '-able' in s]
            name = ', '.join(aff) if aff else next((n for k, n in WF if k in low), s[:40])
        return 'Use of English', 'Словообразование', f'Word formation: {name}'
    if t in ('Collocations / Fixed phrases',):
        return 'Use of English', 'Предлоги и сочетания', 'Idioms & collocations'
    if 'confused' in t.lower():
        return 'Use of English', 'Предлоги и сочетания', 'Words often confused'
    return 'Use of English', 'Предлоги и сочетания', 'Dependent prepositions'


def where(r):
    m = re.search(r'Module (\d+)', r['module'])
    if m:
        return f"М{m.group(1)}" + (f" ({r['lesson']})" if r['lesson'] else '')
    return r['module'].split('·')[0].strip()


STRONG = {'7–9': 'F4B183', '10+': '9DC3E6'}
MID = {'7–9': 'F8CBAD', '10+': 'BDD7EE'}
LIGHT = {'7–9': 'FCE4D6', '10+': 'DDEBF7'}
RANK = {'новое': 3, 'расширение': 2, 'повторение': 1}
SHORT = {'новое': 'новое', 'расширение': 'расш.', 'повторение': 'повт.'}


def load_units():
    """Подтема = будущий урок. Грамматика — из data/subtopics/*.json, Use of English — из строк учебника."""
    rows = sort_rows(load_rows())
    cat_of = {r['topic']: r['category'] for r in rows}
    units = {}
    def unit(block, cat, topic, name, order):
        return units.setdefault(name, {'block': block, 'cat': cat, 'topic': topic, 'order': order, 'cells': {}})
    for f in sorted(glob.glob(os.path.join(D, 'data', 'subtopics', '*.json'))):
        for e in json.load(open(f)):
            t = ALIAS.get(e['topic'], e['topic'])
            u = unit('Грамматика', CATFIX.get(t, cat_of.get(t, '')), t, e['subtopic'].strip(), int(e.get('order') or 99))
            u['cells'].setdefault(int(e['grade']), []).append(
                {'module': e['module'], 'lesson': e.get('lesson') or '', 'status': e['status'], 'detail': e.get('detail') or ''})
    seen = set()
    for r in rows:
        if r['category'] not in LEX:
            continue
        block, cat, name = lex_topic(r)
        topic = name.split(':')[0]
        u = unit(block, cat, topic, name, 1)
        st = 'повторение' if name in seen else 'новое'
        seen.add(name)
        u['cells'].setdefault(r['grade'], []).append(
            {'module': r['module'], 'lesson': r['lesson'], 'status': st, 'detail': r['scope'][:300]})
    return units


def main():
    units = load_units()
    grades = sorted({g for u in units.values() for g in u['cells']})
    videos = {}
    p = os.path.join(D, 'data', 'videos.json')
    if os.path.exists(p):
        videos = json.load(open(p))
    planned, materials = {}, {}
    p = os.path.join(D, 'data', 'materials.json')
    if os.path.exists(p):
        materials = json.load(open(p))
    p = os.path.join(D, 'data', 'planned.json')
    if os.path.exists(p):
        planned = json.load(open(p))

    def status(n, gname):
        have = videos.get(n, {}).get(gname, '')
        if have:
            return 'видео', have
        mat = materials.get(n, {}).get(gname, '')
        plan = planned.get(n, {}).get(gname, '')
        if mat or plan:
            return 'материалы', '\n'.join(x for x in ('материалы: ' + mat.replace('\n', '; ') if mat else '', plan) if x)
        return 'нет', ''

    first = {}
    for n, u in units.items():
        first[u['topic']] = min(first.get(u['topic'], 99), min(u['cells']))
    LEXCATS = ['Фразовые глаголы', 'Словообразование', 'Предлоги и сочетания']
    def key(n):
        u = units[n]
        if u['block'] == 'Грамматика':
            return (0, CATS.index(u['cat']) if u['cat'] in CATS else 99, first[u['topic']], u['topic'], u['order'], n)
        return (1, LEXCATS.index(u['cat']), min(u['cells']), u['topic'], n)
    order = sorted(units, key=key)

    wb = Workbook()
    ws = wb.active; ws.title = 'Как пользоваться'
    for line in [
        'План уроков по грамматике Spotlight',
        '',
        '«Темы × классы»: строка = подтема (один короткий урок: видео → правило → мини-тест). Подтемы сгруппированы по темам: Past Simple делится на правильные/неправильные глаголы, утверждение, отрицание, вопросы, орфографию, каждое значение; сравнения времён — отдельными подтемами.',
        'В клетке класса — где в учебнике (М9 (9b) = Module 9, урок 9b) и статус: «новое» — подтема в Spotlight впервые; «расш.» — была, здесь добавлено; «повт.» — дана снова. Наведите на клетку — во всплывающей подсказке что именно дано (лица, слова-сигналы, список глаголов).',
        'Цвет клетки: оттенок = группа вебинара (оранжевый 7–9 лет, голубой 10+), насыщенность = статус (ярче — новое, бледнее — повторение).',
        'Группы вебинаров: 7–9 лет = Spotlight 2–4 (у Spotlight 2 нет грамматического справочника, поэтому 2 класса нет; 4 класс — в группе 7–9, его темы идут в вебинарах 7–9). 10+ = Spotlight 5–11.',
        'Цвет подтемы: ЗЕЛЁНЫЙ — есть видео; ОРАНЖЕВЫЙ — есть только материалы вебинара (папка на Диске), видео нет — сюда же вебинары из расписания, запись которых ещё не появилась; КРАСНЫЙ — ничего нет. Если подтема нужна обеим группам, цвет — по худшей из двух; по каждой группе — в колонках «Видео 7–9» / «Видео 10+» (видео / материалы / нет), рядом названия видео или папок и дата вебинара. «—» — подтема этой группе не нужна.',
        '«Уроки»: один урок = подтема × группа. Если подтема повторяется в нескольких классах одной группы, это один и тот же урок (Мл-… для 7–9 лет, Ст-… для 10+): в колонке «Где используется» — все классы и модули. Видео подтягивается из «Темы × классы»; «Готовность» отмечается здесь.',
        '«Папки уроков»: Spotlight / класс / модуль / тема (папка) / подтема — какой урок (№) куда положить. Видео и готовность подтягиваются из «Уроков».',
        '«Пересечения»: одна и та же вещь в двух-трёх темах — где предлагаю держать урок, где поставить ссылку.',
        'Утверждение, отрицание и вопрос разделены там, где книга даёт их по отдельности; где справочник даёт их вместе (модальные глаголы старших классов) — одна подтема.',
        '«Кандидаты в вебинары»: по каждой группе — темы с красными подтемами (нет ни видео, ни материалов); отсортировано по классу, где тема появляется. Из них выбираем следующие вебинары.',
        '«Итоги»: сколько уроков сделать каждой группе, у скольких есть видео, сколько готово и опубликовано.',
        '',
        'Источник: грамматический справочник учебников (Use of English — оглавление и приложения). Подробный объём по строкам — в Spotlight_грамматика.xlsx.',
    ]:
        ws.append([line])
    ws['A1'].font = Font(bold=True, size=14)
    ws.column_dimensions['A'].width = 140
    for row in ws.iter_rows(min_row=2):
        row[0].alignment = Alignment(wrap_text=True)

    # --- Темы × классы
    ws = wb.create_sheet('Темы × классы')
    gcols = [f'{g} кл.' for g in grades]
    cols = ['№', 'Блок', 'Категория', 'Тема', 'Подтема (урок)', 'Группы'] + gcols + \
           ['Видео 7–9', 'Вебинар 7–9 (название / ссылка)', 'Видео 10+', 'Вебинар 10+ (название / ссылка)', 'Комментарий']
    header(ws, cols, [5, 12, 16, 22, 46, 8] + [12] * len(grades) + [10, 28, 10, 28, 26])
    ws.freeze_panes = 'F2'
    G0 = 7
    V = G0 + len(grades)
    dv = DataValidation(type='list', formula1='"видео,материалы,нет"', allow_blank=True)
    ws.add_data_validation(dv)
    prev = None
    mrow = {}
    thick = Side(style='medium', color='555555')
    for i, n in enumerate(order, 1):
        u = units[n]
        gs = {group(g) for g in u['cells']}
        grp = 'обе' if len(gs) == 2 else next(iter(gs))
        cells = []
        for g in grades:
            es = u['cells'].get(g, [])
            lines = []
            for e in es:
                l = where(e) + ' · ' + SHORT.get(e['status'], e['status'])
                if l not in lines:
                    lines.append(l)
            cells.append('\n'.join(lines))
        vid = []
        for gname in GROUPS:
            need = grp in (gname, 'обе')
            st, txt = status(n, gname)
            vid += [st if need else '—', txt if need else '']
        ws.append([i, u['block'], u['cat'], u['topic'], n, grp] + cells + vid + [''])
        r = ws.max_row
        mrow[n] = r
        newtopic = u['topic'] != prev
        prev = u['topic']
        for c in ws[r]:
            c.alignment = WRAP
            c.border = Border(left=BOX.left, right=BOX.right, bottom=BOX.bottom, top=thick if newtopic else BOX.top)
        ws.cell(row=r, column=4).font = Font(bold=newtopic, color='000000' if newtopic else '888888')
        sts = [status(n, gname)[0] for gname in GROUPS if grp in (gname, 'обе')]
        ws.cell(row=r, column=5).fill = STFILL[min(sts, key=STRANK.get)]
        for j, g in enumerate(grades):
            es = u['cells'].get(g)
            if not es:
                continue
            c = ws.cell(row=r, column=G0 + j)
            best = max(RANK.get(e['status'], 1) for e in es)
            pal = {3: STRONG, 2: MID, 1: LIGHT}[best]
            c.fill = PatternFill('solid', fgColor=pal[group(g)])
            c.alignment = Alignment(wrap_text=True, vertical='top', horizontal='center')
            if best == 3:
                c.font = Font(bold=True)
            txt = '\n'.join(f"{where(e)}: {e['detail']}" for e in es if e['detail'])
            if txt:
                c.comment = Comment(txt[:1500], 'Spotlight')
                c.comment.width, c.comment.height = 320, 160
        for off in (0, 2):
            c = ws.cell(row=r, column=V + off)
            c.alignment = Alignment(horizontal='center', vertical='top')
            if c.value == '—':
                c.fill = GREY; ws.cell(row=r, column=V + off + 1).fill = GREY
            else:
                dv.add(c)
    last = ws.max_row
    for off in (0, 2):
        col = ws.cell(row=1, column=V + off).column_letter
        rng = f'{col}2:{col}{last}'
        ws.conditional_formatting.add(rng, CellIsRule(operator='equal', formula=['"видео"'], fill=YES))
        ws.conditional_formatting.add(rng, CellIsRule(operator='equal', formula=['"нет"'], fill=NO))
        ws.conditional_formatting.add(rng, CellIsRule(operator='equal', formula=['"материалы"'], fill=SOON))
    ws.auto_filter.ref = ws.dimensions
    vcol = {g: ws.cell(row=1, column=V + 2 * k).column_letter for k, g in enumerate(GROUPS)}
    wcol = {g: ws.cell(row=1, column=V + 2 * k + 1).column_letter for k, g in enumerate(GROUPS)}

    # --- Уроки: один урок на подтему внутри группы; повтор в другом классе той же группы — тот же урок
    mt = "'Темы × классы'"
    PREFIX = {'7–9': 'Мл', '10+': 'Ст'}
    lessons = {}  # (подтема, группа) -> id
    places = {}   # (подтема, группа) -> [(класс, модуль, урок учебника, статус, detail)]
    for gname in GROUPS:
        k = 0
        for n in order:
            gs = sorted(g for g in units[n]['cells'] if group(g) == gname)
            if not gs:
                continue
            k += 1
            lessons[(n, gname)] = f'{PREFIX[gname]}-{k:03d}'
            pl = places.setdefault((n, gname), [])
            for g in gs:
                mods = {}
                for e in units[n]['cells'][g]:
                    m = mods.setdefault(e['module'], {'lessons': [], 'status': e['status'], 'detail': []})
                    if e['lesson'] and e['lesson'] not in m['lessons']:
                        m['lessons'].append(e['lesson'])
                    if RANK.get(e['status'], 1) > RANK.get(m['status'], 1):
                        m['status'] = e['status']
                    if e['detail']:
                        m['detail'].append(e['detail'])
                for m, v in mods.items():
                    pl.append((g, m, ', '.join(v['lessons']), v['status'], '; '.join(v['detail'])))

    wsl = wb.create_sheet('Уроки')
    header(wsl, ['№ урока', 'Группа', 'Блок', 'Тема', 'Подтема (урок)', 'Где используется (класс · модуль · статус)',
                 'Папок', 'Видео', 'Вебинар', 'Готовность', 'Комментарий'],
           [9, 8, 12, 22, 44, 46, 7, 9, 26, 18, 26])
    wsl.freeze_panes = 'F2'
    dvl = DataValidation(type='list', formula1='"не начат,готов без видео,готов с видео,опубликован"', allow_blank=True)
    wsl.add_data_validation(dvl)
    for (n, gname), lid in sorted(lessons.items(), key=lambda x: x[1]):
        u = units[n]
        pl = places[(n, gname)]
        used = '\n'.join(f"{g} кл. · {where({'module': m, 'lesson': les})} · {SHORT.get(st, st)}" for g, m, les, st, _ in pl)
        i = wsl.max_row + 1
        wsl.append([lid, gname, u['block'], u['topic'], n, used, len(pl),
                    f'={mt}!{vcol[gname]}{mrow[n]}',
                    f'={mt}!{wcol[gname]}{mrow[n]}&""',
                    'не начат', ''])
        dvl.add(wsl.cell(row=i, column=10))
    body(wsl)
    for row in wsl.iter_rows(min_row=2):
        row[0].fill = row[1].fill = PatternFill('solid', fgColor=LIGHT[row[1].value])
        row[0].font = row[4].font = Font(bold=True)
        row[4].fill = STFILL[status(row[4].value, row[1].value)[0]]
    for col in ('H',):
        rng = f'{col}2:{col}{wsl.max_row}'
        wsl.conditional_formatting.add(rng, CellIsRule(operator='equal', formula=['"видео"'], fill=YES))
        wsl.conditional_formatting.add(rng, CellIsRule(operator='equal', formula=['"нет"'], fill=NO))
        wsl.conditional_formatting.add(rng, CellIsRule(operator='equal', formula=['"материалы"'], fill=SOON))
    wsl.auto_filter.ref = wsl.dimensions

    # --- Папки уроков: куда кладётся каждый урок
    ws2 = wb.create_sheet('Папки уроков')
    header(ws2, ['Класс', 'Модуль', 'Урок учебника', 'Тема (папка)', 'Подтема (урок)', '№ урока', 'В учебнике',
                 'Что именно в этом классе', 'Группа', 'Видео', 'Готовность', 'Путь'],
           [7, 26, 9, 22, 44, 9, 11, 50, 8, 9, 18, 60])
    ws2.freeze_panes = 'F2'
    items = []
    for (n, gname), pl in places.items():
        for g, m, les, st, det in pl:
            items.append((g, m, les, n, lessons[(n, gname)], st, det))
    items.sort(key=lambda x: (x[0], modnum(x[1]), order.index(x[3])))
    L = "'Уроки'"
    for i, (g, m, les, n, lid, st, det) in enumerate(items, 2):
        det = det[:400]
        if det.startswith('='):
            det = ' ' + det
        ws2.append([g, m, les, units[n]['topic'], n, lid, st, det, group(g),
                    f'=IFERROR(INDEX({L}!$H:$H,MATCH(F{i},{L}!$A:$A,0)),"")',
                    f'=IFERROR(INDEX({L}!$J:$J,MATCH(F{i},{L}!$A:$A,0)),"")',
                    f"Spotlight / {g} класс / {m} / {units[n]['topic']} / {n}"])
    body(ws2)
    for row in ws2.iter_rows(min_row=2):
        grp = row[8].value
        row[0].fill = row[8].fill = PatternFill('solid', fgColor=LIGHT[grp])
        pal = {3: STRONG, 2: MID, 1: LIGHT}[RANK.get(row[6].value, 1)]
        row[6].fill = PatternFill('solid', fgColor=pal[grp])
        row[4].font = Font(bold=True)
        row[4].fill = STFILL[status(row[4].value, row[8].value)[0]]
    rng = f'J2:J{ws2.max_row}'
    ws2.conditional_formatting.add(rng, CellIsRule(operator='equal', formula=['"видео"'], fill=YES))
    ws2.conditional_formatting.add(rng, CellIsRule(operator='equal', formula=['"нет"'], fill=NO))
    ws2.conditional_formatting.add(rng, CellIsRule(operator='equal', formula=['"материалы"'], fill=SOON))
    ws2.auto_filter.ref = ws2.dimensions

    # --- Пересечения
    p = os.path.join(D, 'data', 'overlaps.json')
    if os.path.exists(p):
        wso = wb.create_sheet('Пересечения')
        header(wso, ['№', 'Что повторяется', 'Где встречается (тема: подтема — классы)', 'Где живёт урок (предложение)',
                     'Что делать в остальных местах', 'Решение методиста'], [4, 30, 60, 34, 50, 24])
        wso.freeze_panes = 'C2'
        dvo = DataValidation(type='list', formula1='"согласна,по-другому (см. комментарий)"', allow_blank=True)
        wso.add_data_validation(dvo)
        for i, o in enumerate(json.load(open(p)), 1):
            wso.append([i, o['what'], '\n'.join('• ' + x for x in o['places']), o['home'], o['others'], ''])
            dvo.add(wso.cell(row=wso.max_row, column=6))
        body(wso)
        for row in wso.iter_rows(min_row=2):
            row[1].font = row[3].font = Font(bold=True)

    # --- Кандидаты в вебинары: подтемы без видео и без вебинара в расписании, по темам
    wsk = wb.create_sheet('Кандидаты в вебинары')
    header(wsk, ['Группа', 'Тема', 'С какого класса', 'Подтем без ничего (красные)', 'из них новое в учебнике', 'Красные подтемы', 'Остальные подтемы темы'],
           [8, 30, 10, 10, 12, 70, 40])
    wsk.freeze_panes = 'C2'
    for gname in GROUPS:
        bytopic = {}
        for n in order:
            u = units[n]
            gs = [g for g in u['cells'] if group(g) == gname]
            if not gs:
                continue
            t = bytopic.setdefault((u['block'], u['topic']), {'no': [], 'new': 0, 'have': 0, 'soon': 0, 'first': 99})
            t['first'] = min(t['first'], min(gs))
            st = status(n, gname)[0]
            if st == 'видео':
                t['have'] += 1
            elif st == 'материалы':
                t['soon'] += 1
            else:
                t['no'].append(n.split(': ', 1)[-1])
                if any(e['status'] == 'новое' for g in gs for e in u['cells'][g]):
                    t['new'] += 1
        for (block, topic), t in sorted(bytopic.items(), key=lambda x: (x[1]['first'], -len(x[1]['no']), x[0][1])):
            if not t['no']:
                continue
            wsk.append([gname, topic, f"{t['first']} кл.", len(t['no']), t['new'], '\n'.join('• ' + x for x in t['no']),
                        f"с видео: {t['have']}, только материалы: {t['soon']}" if t['have'] or t['soon'] else ''])
    body(wsk)
    for row in wsk.iter_rows(min_row=2):
        row[0].fill = PatternFill('solid', fgColor=LIGHT[row[0].value])
        row[1].font = Font(bold=True)
    wsk.auto_filter.ref = wsk.dimensions

    # --- Итоги
    ws3 = wb.create_sheet('Итоги')
    header(ws3, ['Группа', 'Классы', 'Уроков сделать', 'Есть видео (зелёные)', 'Только материалы (оранжевые)', 'Ничего нет (красные)', 'Готовы', 'Опубликованы',
                 'Мест в папках (урок × класс × модуль)'],
           [10, 16, 12, 11, 16, 11, 10, 13, 20])
    for gname in GROUPS:
        cls = ', '.join(str(g) for g in grades if group(g) == gname) or '—'
        ws3.append([gname, cls,
                    f'=COUNTIF({L}!B:B,"{gname}")',
                    f'=COUNTIFS({L}!B:B,"{gname}",{L}!H:H,"видео")',
                    f'=COUNTIFS({L}!B:B,"{gname}",{L}!H:H,"материалы")',
                    f'=COUNTIFS({L}!B:B,"{gname}",{L}!H:H,"нет")',
                    f'=COUNTIFS({L}!B:B,"{gname}",{L}!J:J,"готов*")',
                    f'=COUNTIFS({L}!B:B,"{gname}",{L}!J:J,"опубликован")',
                    f"=COUNTIF('Папки уроков'!I:I,\"{gname}\")"])
    body(ws3)
    for row in ws3.iter_rows(min_row=2):
        row[0].fill = PatternFill('solid', fgColor=LIGHT[row[0].value])

    wb.save(OUT)
    print(OUT, len(order), 'подтем;', {g: sum(1 for k in lessons if k[1] == g) for g in GROUPS}, 'уроков;', len(items), 'мест в папках')
    return units, order


if __name__ == '__main__':
    main()
