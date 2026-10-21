"""Вставляет план урока из docs/plan_uroka.md в lesson.html (между /*PLAN_START*/ и /*PLAN_END*/).

python3 splat/tools/build_lesson.py
Запускать после каждой правки плана — урок показывает учителю ровно то, что в файле плана.
"""
import base64, html, io, json, os, re

ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..')


def md(s):
    s = html.escape(s.strip(), quote=False)
    s = re.sub(r'\*\*(.+?)\*\*', r'<b>\1</b>', s)
    s = re.sub(r'\*(.+?)\*', r'<i>\1</i>', s)
    return s


def make_docx(head_lines, rows):
    """Настоящий Word-файл плана: шапка и таблица по разделам. Его же скачивает учитель из урока."""
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.shared import Pt, Cm, RGBColor
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = sec.page_height, sec.page_width
    for side in ('left_margin', 'right_margin', 'top_margin', 'bottom_margin'):
        setattr(sec, side, Cm(1.5))
    doc.styles['Normal'].font.name = 'Calibri'
    doc.styles['Normal'].font.size = Pt(10.5)
    h = doc.add_heading('План урока «Пингвинёнок идёт домой» (Splat Kids × WowSpeak)', 1)
    def runs(par, text):
        for k, part in enumerate(re.split(r'\*\*(.+?)\*\*', text)):
            for j, sub in enumerate(re.split(r'\*(.+?)\*', part)):
                r = par.add_run(sub); r.bold = k % 2 == 1; r.italic = j % 2 == 1
    for line in head_lines:
        runs(doc.add_paragraph(), line)
    cur, t = None, None
    for r in rows:
        if r['sec'] != cur:
            cur = r['sec']
            doc.add_heading(cur, 2)
            t = doc.add_table(rows=1, cols=4); t.style = 'Table Grid'
            for c, name in zip(t.rows[0].cells, ['Слайд', 'Что нужно сделать', "Teacher's language", 'Suggested answers']):
                c.text = ''; rr = c.paragraphs[0].add_run(name); rr.bold = True
        cells = t.add_row().cells
        for c, val in zip(cells, [str(r['n']), r['todo_md'], r['say_md'], r['ans_md']]):
            c.text = ''; runs(c.paragraphs[0], val)
    for tbl in doc.tables:
        for row in tbl.rows:
            for c, w in zip(row.cells, (1.4, 6.5, 12.5, 5.5)):
                c.width = Cm(w)
    buf = io.BytesIO(); doc.save(buf); return buf.getvalue()


def main():
    rows, sec = [], ''
    for line in open(os.path.join(ROOT, 'docs', 'plan_uroka.md'), encoding='utf-8'):
        if line.startswith('## '):
            sec = line[3:].strip()
        m = re.match(r'^\|\s*(\d+)\s*\|(.*)\|\s*$', line)
        if m:
            cols = [c for c in m.group(2).split('|')]
            todo, say, ans = (cols + ['', '', ''])[:3]
            rows.append({'n': int(m.group(1)), 'sec': md(sec), 'todo': md(todo), 'say': md(say), 'ans': md(ans),
                         'sec_md': sec, 'todo_md': todo.strip(), 'say_md': say.strip(), 'ans_md': ans.strip()})
    src_head = [l for l in open(os.path.join(ROOT, 'docs', 'plan_uroka.md'), encoding='utf-8').read().split('---')[0].splitlines()[1:] if l.strip()]
    docx = make_docx(src_head, [dict(r, sec=r['sec_md']) for r in rows])
    open(os.path.join(ROOT, 'docs', 'plan_uroka.docx'), 'wb').write(docx)
    for r in rows:
        for k in ('sec_md', 'todo_md', 'say_md', 'ans_md'):
            r.pop(k)
    p = os.path.join(ROOT, 'lesson.html')
    s = open(p, encoding='utf-8').read()
    # шапка плана (всё до первой таблицы) — для скачивания плана целиком
    src = open(os.path.join(ROOT, 'docs', 'plan_uroka.md'), encoding='utf-8').read()
    head = [md(l) for l in src.split('---')[0].splitlines()[1:] if l.strip()]
    s = re.sub(r'/\*HEAD_START\*/.*?/\*HEAD_END\*/',
               lambda _: '/*HEAD_START*/' + json.dumps(head, ensure_ascii=False) + '/*HEAD_END*/', s, flags=re.S)
    s = re.sub(r'/\*PLAN_START\*/.*?/\*PLAN_END\*/',
               lambda _: '/*PLAN_START*/' + json.dumps(rows, ensure_ascii=False) + '/*PLAN_END*/', s, flags=re.S)
    s = re.sub(r'/\*DOCX\*/.*?/\*DOCX_END\*/', lambda _: '/*DOCX*/' + json.dumps(base64.b64encode(docx).decode()) + '/*DOCX_END*/', s, flags=re.S)
    open(p, 'w', encoding='utf-8').write(s)
    print('слайдов в плане:', len(rows))


if __name__ == '__main__':
    main()
